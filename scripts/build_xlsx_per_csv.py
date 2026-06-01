"""Update SYSHECF H2 CSVs + build SHELF/SYSHECF xlsx with one tab per CSV.

For every SHELF and SYSHECF folder in scope (national + 48 states):

1. SYSHECF-hydrogen-CT.csv  ← copy DATA rows from SYSHECF-natural-gas-peaker.csv
                              (preserve hydrogen-CT's first-cell label)
2. SYSHECF-hydrogen-CC.csv  ← copy DATA rows from SYSHECF-combined-cycle.csv
                              (preserve hydrogen-CC's first-cell label)
3. Build `<location label> - Hourly Demand and SHELF.xlsx` — one tab per
   SHELF-*.csv (excluding SHELF-days-per-timeslice which gets its own tab),
   each tab populated exactly like the CSV.
4. Build `<location label> - SYSHECF.xlsx` — one tab per SYSHECF-*.csv.

Tab names use the CSV stem (e.g. "SHELF-residential-cooling",
"SYSHECF-onshore-wind"). Excel sheet name limit is 31 chars — all names fit.

Run:
    python scripts/build_xlsx_per_csv.py
"""
from __future__ import annotations
import csv
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent.parent
REPO = Path(r"C:\Users\RobbieOrvis\Models\state-eps-data-repository")
EPS_US_SHELF   = Path(r"C:\Users\RobbieOrvis\Models\US\Models\eps-us\InputData\elec\SHELF")
EPS_US_SYSHECF = Path(r"C:\Users\RobbieOrvis\Models\US\Models\eps-us\InputData\elec\SYSHECF")


def _read_csv_rows(path: Path) -> list[list[str]]:
    with open(path, newline='') as f:
        return list(csv.reader(f))


def _write_csv_rows(path: Path, rows: list[list[str]]):
    with open(path, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerows(rows)


def copy_data_rows_into(target_csv: Path, source_csv: Path):
    """Overwrite the slice data rows of target_csv with source_csv's values,
    keeping target's first cell label (e.g., "hydrogen CT") on the header row."""
    if not (target_csv.exists() and source_csv.exists()):
        return False
    target = _read_csv_rows(target_csv)
    source = _read_csv_rows(source_csv)
    if not target or not source:
        return False
    # Header row: target keeps its first cell (e.g., 'hydrogen CT'); rest copied from source
    new_rows = [[target[0][0]] + source[0][1:]]
    # Data rows: copy source's slice rows verbatim
    new_rows.extend(source[1:])
    _write_csv_rows(target_csv, new_rows)
    return True


def update_hydrogen_csvs(syshecf_dir: Path) -> tuple[bool, bool]:
    ct = copy_data_rows_into(
        syshecf_dir / 'SYSHECF-hydrogen-CT.csv',
        syshecf_dir / 'SYSHECF-natural-gas-peaker.csv',
    )
    cc = copy_data_rows_into(
        syshecf_dir / 'SYSHECF-hydrogen-CC.csv',
        syshecf_dir / 'SYSHECF-combined-cycle.csv',
    )
    return ct, cc


def build_xlsx_from_csv_folder(folder: Path, out_xlsx: Path, prefix: str):
    """For every <prefix>-*.csv in folder, create a tab in out_xlsx with the
    CSV's contents. Tab name = CSV stem."""
    files = sorted(folder.glob(f'{prefix}-*.csv'))
    if not files:
        return False
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    for csvp in files:
        tab_name = csvp.stem  # e.g., 'SHELF-residential-cooling'
        # Excel sheet name limits: 31 chars; certain chars disallowed
        tab_name = tab_name[:31]
        ws = wb.create_sheet(title=tab_name)
        rows = _read_csv_rows(csvp)
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                # Try numeric conversion for data cells (Hour0..Hour23 values
                # in data rows). Header text stays as-is.
                if r_idx >= 2 and c_idx >= 2:
                    try:
                        ws.cell(row=r_idx, column=c_idx, value=float(val))
                        continue
                    except (TypeError, ValueError):
                        pass
                ws.cell(row=r_idx, column=c_idx, value=val)
        # Set reasonable column widths
        ws.column_dimensions['A'].width = max(20, min(40, len(rows[0][0]) + 2))
        for c in range(2, 26):
            ws.column_dimensions[get_column_letter(c)].width = 10
    out_xlsx.parent.mkdir(exist_ok=True, parents=True)
    try:
        wb.save(out_xlsx)
    except PermissionError:
        print(f"    WARN: cannot write {out_xlsx} (file open in Excel?)")
        return False
    return True


def process_location(label: str, shelf_dir: Path, syshecf_dir: Path,
                     shelf_xlsx: Path, syshecf_xlsx: Path) -> dict:
    info = {'label': label, 'shelf_dir': str(shelf_dir), 'syshecf_dir': str(syshecf_dir),
            'h2_ct': False, 'h2_cc': False, 'shelf_xlsx': False, 'syshecf_xlsx': False}
    if syshecf_dir.exists():
        info['h2_ct'], info['h2_cc'] = update_hydrogen_csvs(syshecf_dir)
        info['syshecf_xlsx'] = build_xlsx_from_csv_folder(syshecf_dir, syshecf_xlsx, 'SYSHECF')
    if shelf_dir.exists():
        info['shelf_xlsx'] = build_xlsx_from_csv_folder(shelf_dir, shelf_xlsx, 'SHELF')
    return info


def main():
    results = []

    # National (eps-us) — use canonical Vensim filenames
    info = process_location(
        'US national (eps-us)',
        EPS_US_SHELF,
        EPS_US_SYSHECF,
        EPS_US_SHELF / 'Seasonal Hourly Equipment Load Factors by End Use.xlsx',
        EPS_US_SYSHECF / 'Start Year Seasonal Expected Hourly Electricity Capacity Factors.xlsx',
    )
    results.append(info)
    print(f"[national] H2-CT updated: {info['h2_ct']}; H2-CC updated: {info['h2_cc']}; "
          f"SHELF xlsx: {info['shelf_xlsx']}; SYSHECF xlsx: {info['syshecf_xlsx']}")

    # All 48 lower-48 states
    states = sorted([p.name for p in REPO.iterdir() if p.is_dir() and len(p.name) == 2])
    states = [s for s in states if s not in {'AK', 'HI', 'DC'}]
    for state in states:
        shelf_dir = REPO / state / 'elec' / 'SHELF' / '_python_pipeline'
        syshecf_dir = REPO / state / 'elec' / 'SYSHECF' / '_python_pipeline'
        shelf_xlsx = shelf_dir / f'{state} - Hourly Demand and SHELF.xlsx'
        syshecf_xlsx = syshecf_dir / f'{state} - SYSHECF.xlsx'
        info = process_location(state, shelf_dir, syshecf_dir, shelf_xlsx, syshecf_xlsx)
        results.append(info)
        msg = f"[{state}] H2-CT: {info['h2_ct']}; H2-CC: {info['h2_cc']}; " \
              f"SHELF xlsx: {info['shelf_xlsx']}; SYSHECF xlsx: {info['syshecf_xlsx']}"
        print(msg)

    # Summary
    n_h2_ct = sum(1 for r in results if r['h2_ct'])
    n_h2_cc = sum(1 for r in results if r['h2_cc'])
    n_shelf = sum(1 for r in results if r['shelf_xlsx'])
    n_syshecf = sum(1 for r in results if r['syshecf_xlsx'])
    print(f"\nSummary across {len(results)} locations:")
    print(f"  H2-CT CSV updated:     {n_h2_ct}")
    print(f"  H2-CC CSV updated:     {n_h2_cc}")
    print(f"  SHELF xlsx written:    {n_shelf}")
    print(f"  SYSHECF xlsx written:  {n_syshecf}")


if __name__ == '__main__':
    main()
