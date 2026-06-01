"""Fetch per-state capacity factors from EIA State Electricity Profiles.

Downloads each state's `SEP Tables for <STATE>.xlsx` and extracts the
2024 annual capacity factors from Sheet 15 ('15 Capacity Factors Annual')
plus distributed solar CF from Sheet 19 ('19. Small Scale Solar Annual').

Output: `data/eia_state_cfs.csv`
Columns: state_iso2, tech, cf_target_2024, capacity_mw_2024
Tech names use SYSHECF naming convention (solar-pv, onshore-wind, etc.).

Source: EIA State Electricity Profiles
URL pattern: https://www.eia.gov/electricity/state/xls/SEP%20Tables%20for%20<STATE>.xlsx

Run:
    python scripts/fetch_eia_state_cfs.py

Notes:
- States where a tech has zero/missing capacity are excluded from that tech's row.
- Distributed solar (solar-pv-dist) CF is computed as
  generation_2024 / (capacity_2024 * 8760) since Sheet 19 doesn't pre-compute it.
- All-zero or blank CF cells are skipped.
"""
from __future__ import annotations
import csv
import time
import urllib.request
import urllib.parse
from pathlib import Path

import openpyxl

ROOT = Path(__file__).parent.parent
PROFILES_DIR = ROOT / "data" / "eia_state_profiles"
OUTPUT_CSV = ROOT / "data" / "eia_state_cfs.csv"

PROFILES_DIR.mkdir(parents=True, exist_ok=True)

# All 50 states + DC. AK and HI included (some states may lack certain techs).
STATES = [
    'AL', 'AK', 'AZ', 'AR', 'CA', 'CO', 'CT', 'DE', 'DC', 'FL', 'GA', 'HI',
    'ID', 'IL', 'IN', 'IA', 'KS', 'KY', 'LA', 'ME', 'MD', 'MA', 'MI', 'MN',
    'MS', 'MO', 'MT', 'NE', 'NV', 'NH', 'NJ', 'NM', 'NY', 'NC', 'ND', 'OH',
    'OK', 'OR', 'PA', 'RI', 'SC', 'SD', 'TN', 'TX', 'UT', 'VT', 'VA', 'WA',
    'WV', 'WI', 'WY',
]

# Map EIA Sheet 15 tech names → SYSHECF tech names
# EIA reports CFs as percent (e.g., 23.4 means 23.4%). We divide by 100.
EIA_TO_SYSHECF = {
    'Nuclear': 'nuclear',
    'Natural Gas - CC': 'combined-cycle',
    'Natural Gas - GT': 'natural-gas-peaker',
    'Natural Gas - ST': 'steam-turbine',
    'Coal': 'hard-coal',
    'Wind - Onshore': 'onshore-wind',
    'Wind - Offshore': 'offshore-wind',
    'Solar - PV': 'solar-pv',
    'Solar - Thermal': 'solar-thermal',
    'Hydroelectric': 'hydro',
    'Wood': 'biomass',          # primary biomass type
    'Geothermal': 'geothermal',
    'Petroleum - ST': 'petroleum',
    'Petroleum - GT': 'petroleum',  # combined later
    'Petroleum - IC': 'petroleum',
}

# Distributed solar from Sheet 19: Total generation row 6, capacity row 21
# (capacity row index may vary; we search by label)
HEADERS = {"User-Agent": "Mozilla/5.0 (Energy Innovation EPS calibration)"}


def _url(state: str) -> str:
    name = f"SEP Tables for {state}.xlsx"
    return f"https://www.eia.gov/electricity/state/xls/{urllib.parse.quote(name)}"


def download_state(state: str, force: bool = False) -> Path:
    path = PROFILES_DIR / f"{state}.xlsx"
    if path.exists() and not force and path.stat().st_size > 1000:
        return path
    url = _url(state)
    req = urllib.request.Request(url, headers=HEADERS)
    with urllib.request.urlopen(req, timeout=60) as resp:
        data = resp.read()
    path.write_bytes(data)
    return path


def parse_sheet_15(path: Path) -> dict:
    """Return dict: tech_syshecf -> {'cf': float, 'capacity_mw': float} for techs present."""
    wb = openpyxl.load_workbook(path, data_only=True)
    if '15 Capacity Factors Annual' not in wb.sheetnames:
        return {}
    ws = wb['15 Capacity Factors Annual']
    # Row 3 is header with year columns. "2024 Time adjusted capacity (MW)" is column B (index 2).
    # "Year 2024" is column C (index 3).
    out: dict[str, dict] = {}
    pet_total_cap = 0.0
    pet_total_gen = 0.0  # weighted-by-cap CF computed at end
    pet_count = 0
    for r in ws.iter_rows(min_row=5, max_row=30, values_only=True):
        if not r or not r[0]:
            continue
        label = str(r[0]).strip()
        if label not in EIA_TO_SYSHECF:
            continue
        syshecf_name = EIA_TO_SYSHECF[label]
        cap = r[1] if r[1] not in (None, '', ' ') else None
        cf_pct = r[2] if r[2] not in (None, '', ' ') else None
        if cap is None or cf_pct is None:
            continue
        try:
            cap_f = float(cap)
            cf = float(cf_pct) / 100.0
        except (TypeError, ValueError):
            continue
        if cap_f <= 0 or cf <= 0 or cf > 1.0:
            continue
        # For petroleum: weighted aggregation of GT/ST/IC by capacity
        if syshecf_name == 'petroleum':
            pet_total_cap += cap_f
            pet_total_gen += cap_f * cf  # CF-weighted by cap
            pet_count += 1
            continue
        # For biomass: Wood is primary; "Other, Biomass" could be added but Wood is the canonical proxy
        if syshecf_name in out:
            # Combine by capacity-weighted CF
            prev = out[syshecf_name]
            tot_cap = prev['capacity_mw'] + cap_f
            tot_gen = prev['capacity_mw'] * prev['cf'] + cap_f * cf
            out[syshecf_name] = {'cf': tot_gen / tot_cap, 'capacity_mw': tot_cap}
        else:
            out[syshecf_name] = {'cf': cf, 'capacity_mw': cap_f}
    if pet_total_cap > 0:
        out['petroleum'] = {'cf': pet_total_gen / pet_total_cap, 'capacity_mw': pet_total_cap}
    return out


def parse_sheet_19_distpv(path: Path) -> dict | None:
    """Parse small-scale (distributed) solar from Sheet 19.

    Returns {'cf': float, 'capacity_mw': float} or None if not available.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    if '19. Small Scale Solar Annual' not in wb.sheetnames:
        return None
    ws = wb['19. Small Scale Solar Annual']
    # Need: 2024 Total Generation (MWh) and 2024 Total Capacity (MW)
    # Sheet has subsections: Generation (MWh) row, Capacity (MW) row, etc.
    # Look for two 'Total' rows; the first is under "Generation (MWh)" header, the
    # second is under "Capacity (MW)" header.
    gen_total = None
    cap_total = None
    section = None
    for r in ws.iter_rows(min_row=1, max_row=40, values_only=True):
        if not r:
            continue
        first = str(r[0]).strip() if r[0] else ''
        # Detect section headers
        if 'Generation (MWh)' in first:
            section = 'gen'
            continue
        if 'Capacity (MW)' in first:
            section = 'cap'
            continue
        if first == 'Total' and section == 'gen' and gen_total is None:
            gen_total = r[1] if r[1] not in (None, '', ' ') else None
        elif first == 'Total' and section == 'cap' and cap_total is None:
            cap_total = r[1] if r[1] not in (None, '', ' ') else None
    if gen_total is None or cap_total is None:
        return None
    try:
        gen_mwh = float(gen_total)
        cap_mw = float(cap_total)
        if cap_mw <= 0:
            return None
        cf = gen_mwh / (cap_mw * 8760.0)
        if cf <= 0 or cf > 0.5:  # sanity bound for distpv CF
            return None
        return {'cf': cf, 'capacity_mw': cap_mw}
    except (TypeError, ValueError):
        return None


def main():
    rows = [['state_iso2', 'tech', 'cf_target_2024', 'capacity_mw_2024']]
    print(f"Fetching EIA State Electricity Profiles for {len(STATES)} states...")
    for i, state in enumerate(STATES, 1):
        try:
            path = download_state(state)
            sheet15 = parse_sheet_15(path)
            sheet19 = parse_sheet_19_distpv(path)
            if sheet19:
                sheet15['solar-pv-dist'] = sheet19
            n_techs = len(sheet15)
            print(f"  [{i:>2}/{len(STATES)}] {state}: {n_techs} techs parsed")
            for tech in sorted(sheet15.keys()):
                d = sheet15[tech]
                rows.append([state, tech, f"{d['cf']:.4f}", f"{d['capacity_mw']:.1f}"])
            time.sleep(0.3)  # be polite to EIA
        except Exception as e:
            print(f"  [{i:>2}/{len(STATES)}] {state}: FAILED ({e})")
    with open(OUTPUT_CSV, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerows(rows)
    print(f"\nWrote {len(rows)-1} (state, tech) rows to {OUTPUT_CSV}")


if __name__ == '__main__':
    main()
