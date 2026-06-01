"""H3: Inspect the VA SHELF workbook to learn how it was built."""
import os
from openpyxl import load_workbook

WB = r"C:\Users\RobbieOrvis\Models\state-eps-data-repository\VA\elec\SHELF\Seasonal Hourly Equipment Load Factors by End Use.xlsx"

wb = load_workbook(WB, data_only=False, read_only=True)
print("Sheets:")
for s in wb.sheetnames:
    ws = wb[s]
    print(f"  {s}: max_row={ws.max_row} max_col={ws.max_column}")
